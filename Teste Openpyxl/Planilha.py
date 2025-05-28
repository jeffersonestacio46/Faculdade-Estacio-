import random
from openpyxl import Workbook

workbook = Workbook()

sheet = workbook.active
sheet.title = 'Estoque'
#A baixo está a linha 1 da planilha, onde os headers serão escrito como cabeçalhos
headers = ['Nome do produto', 'Valor do fornecedor', 'Lucratividade(%)', 'Quantidade']

for col_num, header in enumerate(headers, start=1):
    sheet.cell(row=1, column=col_num, value=header)
#Função de gerador de nomes de produtos
def gerar_nome_produto():
    prefixos = ['Super', 'ultra', 'Power', 'Eco', 'Max']#Para criação de lista usa o colchetes
    tipos = ['widget', 'Gadget', 'Device', 'Tool', 'Instrument', 'Appliance']
    sufixos = ['Plus', 'Pro', 'X', '2000', 'Prime', 'Elite']
    #Criando uma interação de forma aletoria(random)
    return f'{random.choice(prefixos)}{random.choice(tipos)}{random.choice(sufixos)}'

num_produto = 50 #Informando o sistema que quero que ele gere 50 produtos para mim

for row_num in range(2, num_produto + 2):
    nome_produto = gerar_nome_produto()
   #Uso o round para arendodar o numero/random.uniform digo o renged de numero que irá criar, e no segundo paramento para ter no maximo duas casas decimais
    valor_fornecedor = round(random.uniform(10.0, 500.0),2)
    lucratividade = random.randint(10, 100)#random.randint gera um numero inteiro aleatorio(Nesse caso mostra o lucro de 10 até 100%, porém não está em pocentagem)
    quantidade = random.randint(1, 100)
    #Inserindo os dados acima na planilha com os codigos a baixo
    sheet.cell(row=row_num, column=1, value=nome_produto)
    sheet.cell(row=row_num, column=2, value=valor_fornecedor)
    sheet.cell(row=row_num, column=3, value=lucratividade)
    sheet.cell(row=row_num, column=4, value=quantidade)

file_path = 'estoque.xlsx'
workbook.save(file_path)

