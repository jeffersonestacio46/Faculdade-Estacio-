from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

file_path = 'estoque.xlsx'

try:
    workbook = load_workbook(file_path)
    sheet = workbook['Estoque']
except FileNotFoundError:
    print(f"Erro: O arquivo '{file_path}' não foi encontrado.")
    exit()
except KeyError:
    print(f"Erro: A planilha 'Estoque' não foi encontrada no arquivo.")
    exit()

col_nome_produto = 1  # Coluna A
col_valor_fornecedor = 2  # Coluna B
col_lucratividade = 3  # Coluna C
col_quantidade = 4  # Coluna D

col_preco_venda = col_quantidade + 1  # Adicionar ao lado dela (Coluna E)
col_lucro_total = col_preco_venda + 1  # Coluna F
col_valor_total = col_lucro_total + 1  # Coluna G

# Fazer os novos cabeçalhos
sheet.cell(row=1, column=col_preco_venda, value='Preço de venda')
sheet.cell(row=1, column=col_lucro_total, value='Lucro Total')
sheet.cell(row=1, column=col_valor_total, value='Valor total')

max_row = sheet.max_row

# Fórmulas
for row in range(2, max_row + 1):
    cell_valor_fornecedor = sheet.cell(row=row, column=col_valor_fornecedor)
    cell_lucratividade = sheet.cell(row=row, column=col_lucratividade)
    cell_quantidade = sheet.cell(row=row, column=col_quantidade)
    cell_preco_venda_obj = sheet.cell(row=row, column=col_preco_venda)
    cell_lucro_total_obj = sheet.cell(row=row, column=col_lucro_total)
    cell_valor_total_obj = sheet.cell(row=row, column=col_valor_total)

    formula_preco_venda = f'={cell_valor_fornecedor.coordinate}*(1+{cell_lucratividade.coordinate}/100)'
    cell_preco_venda_obj.value = formula_preco_venda

    formula_lucro_total = f'=({cell_preco_venda_obj.coordinate}-{cell_valor_fornecedor.coordinate})*{cell_quantidade.coordinate}'
    cell_lucro_total_obj.value = formula_lucro_total

    formula_valor_total = f'={cell_preco_venda_obj.coordinate}*{cell_quantidade.coordinate}'
    cell_valor_total_obj.value = formula_valor_total

linha_total = max_row + 2
sheet.cell(row=linha_total, column=col_nome_produto, value='Totais Gerais')
sheet.merge_cells(start_row=linha_total, start_column=col_nome_produto, end_row=linha_total, end_column=col_quantidade)

# Fórmulas de total
formula_total_lucro = f'=SUM({get_column_letter(col_lucro_total)}2:{get_column_letter(col_lucro_total)}{max_row})'
sheet.cell(row=linha_total, column=col_lucro_total).value = formula_total_lucro

formula_total_valor = f'=SUM({get_column_letter(col_valor_total)}2:{get_column_letter(col_valor_total)}{max_row})'
sheet.cell(row=linha_total, column=col_valor_total).value = formula_total_valor

try:
    workbook.save(file_path)
    print(f"Arquivo '{file_path}' salvo com sucesso!")
except Exception as e:
    print(f"Erro ao salvar o arquivo: {e}")