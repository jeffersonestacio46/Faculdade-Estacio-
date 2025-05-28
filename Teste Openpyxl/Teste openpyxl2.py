import openpyxl
# Criar uma nova planilha (workbook)
book = openpyxl.Workbook()
# Visualizar páginas existentes (inicialmente, haverá apenas uma página padrão)
print(book.sheetnames)
# Criar uma nova página chamada 'Frutas'
frutas_page = book.create_sheet('Frutas')
# Adicionar dados à página 'Frutas'
frutas_page.append(['Nome', 'Quantidade', 'Preço'])  # Cabeçalho
frutas_page.append(['Banana', 5, 'R$3,90'])
frutas_page.append(['Fruta 2', 2, 'R$15,90'])
frutas_page.append(['Fruta 3', 10, 'R$30,90'])
frutas_page.append(['Fruta 4', 2, 'R$50,50'])
# Salvar a planilha
book.save('Planilha de Compras.xlsx')