import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill #Manipular fonte e preenchimento, o openpyxl.styles manipula os estilos

from gerador_de_dados import df #Importou o arquivo python e importou o df do codigo
#O ExcelWriter escreve a planilha usando o openpyxl, depois colocamos o engine=opnepyxl, e o as com uma variavel
with pd.ExcelWriter('relatorio_vendas.xlsx', engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name='Vendas', index=False)#O to_excel fala que vamos escrever para um excel
    ws = writer.sheets['Vendas']#ws é a variavel que temos de referencia da planilha
    
    header_fill = PatternFill(start_color='FFFF00', end_color='FFFF00',
    fill_type='solid')#Cor amarela solida
    
    for cell in ws[1]:
        cell.font = Font(bold=True)#Aplicamos um fonte
        cell.fill = header_fill#Aplicamos um preenchimento