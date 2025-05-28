import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Código para gerar os dados
data = {
    'Produtos': ['Teclado', 'Mouse', 'Monitor', 'Notebook', 'Impressora'],
    'Quantidade': [10, 15, 7, 5, 8],
    'Preço Unitario': [120.00, 80.00, 750.00, 3200.00, 900.00]
}
df = pd.DataFrame(data)
df['Total'] = df['Quantidade'] * df['Preço Unitario']

# Código para gerar a planilha Excel
with pd.ExcelWriter('relatorio_vendas.xlsx', engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name='Vendas', index=False)
    ws = writer.sheets['Vendas']

    header_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    # Ajusta a largura das colunas automaticamente
    for i, column_cells in enumerate(ws.columns, 1):
        length = max(len(str(cell.value)) for cell in column_cells)
        ws.column_dimensions[get_column_letter(i)].width = length + 2

print("Relatório de vendas gerado com sucesso!")

